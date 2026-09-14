"""Casos de uso administrativos sin acceso directo a PostgreSQL."""

from __future__ import annotations

import hashlib
import json
from io import BytesIO
from typing import Any
from uuid import NAMESPACE_URL, UUID, uuid4, uuid5

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from ...core.security import hash_password
from ..evaluaciones.rules import normalize_evaluation
from ..evaluaciones.schemas import EvaluationCreate, ModuleKey
from .repository import (
    AdminRepository,
    EvaluationNotFoundError,
    ImportPreviewError,
)
from .schemas import (
    AdminEvaluationCorrection,
    AdminEvaluationDetail,
    AdminEvaluationPage,
    AdminEvaluationQuery,
    AdminEvaluationSummary,
    AdminLoadPage,
    AdminLoadQuery,
    AdminMasterPage,
    AdminMasterQuery,
    AdminQABulkResolution,
    AdminQAPage,
    AdminQAQuery,
    AdminQAReview,
    AdminQASummary,
    EvaluationCorrectionRequest,
    ImportManifest,
    ImportPreview,
    ImportPreviewRow,
    ImportResult,
    MasterMutationRequest,
    MasterResource,
    QABulkDuplicateRequest,
    QAReviewRequest,
    UserMutationRequest,
)
from .analytics import AnalyticsQuery, EvaluationTrend, EvaluationAnalytics

TEMPLATE_HEADERS = [
    "client_id",
    "module_key",
    "fecha",
    "captured_at",
    "lote_id",
    "fundo",
    "modulo",
    "lote",
    "cortina",
    "hilera",
    "planta",
    "evaluador_dni",
    "valores_json",
]
TEMPLATE_EXAMPLE = [
    "",
    "estadios",
    "2026-09-02",
    "2026-09-02T10:30:00Z",
    "12",
    "",
    "",
    "",
    "1",
    "2",
    "3",
    "10616663",
    '{"m1_e1": 1, "m1_e2": 2}',
]


class AdminService:
    def __init__(self, repository: AdminRepository):
        self._repository = repository

    def list_evaluations(self, query: AdminEvaluationQuery) -> AdminEvaluationPage:
        return self._repository.list_evaluations(query)

    def evaluation_counts(self, query):
        return self._repository.evaluation_counts(query)

    def evaluation_summary(self, query: AdminEvaluationQuery) -> AdminEvaluationSummary:
        return self._repository.evaluation_summary(query)

    def evaluation_trend(self, query: AnalyticsQuery) -> EvaluationTrend:
        return self._repository.evaluation_trend(query)

    def evaluation_analytics(self, query: AnalyticsQuery) -> EvaluationAnalytics:
        return self._repository.evaluation_analytics(query)

    def export_evaluations(self, query: AdminEvaluationQuery):
        return self._repository.export_evaluations(query)

    def get_evaluation(
        self, module_key: ModuleKey, source_id: int, source_table: str | None = None
    ) -> AdminEvaluationDetail | None:
        if source_table is None:
            return self._repository.get_evaluation(module_key, source_id)
        return self._repository.get_evaluation(module_key, source_id, source_table)

    def correct_evaluation(
        self,
        module_key: ModuleKey,
        source_id: int,
        request: EvaluationCorrectionRequest,
        usuario_id: int,
    ) -> AdminEvaluationCorrection:
        detail = self._repository.get_evaluation(module_key, source_id, request.source_table)
        if detail is None:
            raise EvaluationNotFoundError("La evaluación no existe o está fuera de su alcance")
        if detail.module_key != module_key or detail.source_table != request.source_table:
            raise EvaluationNotFoundError("La tabla de origen no coincide con la evaluación")
        if (
            detail.source_table == "ev_evaluacion"
            and detail.detalle.get("publicacion_id") is not None
        ):
            raise ValueError(
                "Histórico verificado: requiere conciliación; "
                "no se sobrescribe la publicación raw-core"
            )
        values = _correction_values(module_key, source_id, request)
        mutation = self._repository.correct_evaluation(
            module_key, source_id, request, values, usuario_id
        )
        updated = self._repository.get_evaluation(module_key, source_id, request.source_table)
        if updated is None:
            raise EvaluationNotFoundError("La evaluación corregida no pudo ser consultada")
        return AdminEvaluationCorrection(mutation=mutation, evaluacion=updated)

    def list_master(self, resource: MasterResource, query: AdminMasterQuery) -> AdminMasterPage:
        return self._repository.list_master(resource, query)

    def mutate_master(
        self,
        resource: MasterResource,
        resource_id: int | None,
        request: MasterMutationRequest,
        usuario_id: int,
    ) -> Any:
        if not request.valores:
            raise ValueError("La mutación del maestro requiere valores")
        return self._repository.mutate_master(resource, resource_id, request, usuario_id)

    def list_roles(self, query: AdminMasterQuery) -> AdminMasterPage:
        return self._repository.list_roles(query)

    def mutate_user(
        self,
        usuario_id: int | None,
        request: UserMutationRequest,
        actor_id: int,
    ) -> Any:
        if usuario_id is None and (not request.email or not request.nombre or not request.rol):
            raise ValueError("Un usuario nuevo requiere email, nombre y rol")
        values = request.model_dump(
            exclude={"idempotency_key", "comentario", "password", "alcances"},
            exclude_none=True,
            mode="json",
        )
        if request.password is not None:
            values["hash_password"] = hash_password(request.password)
        if request.alcances is not None:
            values["alcances"] = [
                scope.model_dump(exclude_none=True) for scope in request.alcances
            ]
        if not values:
            raise ValueError("La mutación del usuario requiere al menos un campo")
        return self._repository.mutate_user(usuario_id, request, actor_id, values)

    def list_users(self, query: AdminMasterQuery) -> AdminMasterPage:
        return self._repository.list_users(query)

    def list_imports(self, query: AdminLoadQuery) -> AdminLoadPage:
        return self._repository.list_imports(query)

    def qa_summary(self) -> AdminQASummary:
        return self._repository.qa_summary()

    def list_qa(self, query: AdminQAQuery) -> AdminQAPage:
        return self._repository.list_qa(query)

    def review_qa(
        self, rechazo_id: int, request: QAReviewRequest, usuario_id: int
    ) -> AdminQAReview:
        return self._repository.review_qa(rechazo_id, request, usuario_id)

    def confirm_qa_duplicates(
        self, request: QABulkDuplicateRequest, usuario_id: int
    ) -> AdminQABulkResolution:
        return self._repository.confirm_qa_duplicates(request, usuario_id)

    def create_import_preview(
        self, content: bytes, filename: str, usuario_id: int
    ) -> ImportPreview:
        preview = self.preview_xlsx(content, filename)
        return self._repository.create_import_preview(preview, usuario_id)

    def claim_import(
        self, preview_id: UUID, file_sha256: str, usuario_id: int
    ) -> ImportManifest:
        return self._repository.claim_import(str(preview_id), file_sha256, usuario_id)

    def complete_import(self, preview_id: UUID, usuario_id: int, result: ImportResult) -> None:
        self._repository.complete_import(str(preview_id), usuario_id, result)

    def fail_import(self, preview_id: UUID, usuario_id: int, error: str) -> None:
        self._repository.fail_import(str(preview_id), usuario_id, error)

    @staticmethod
    def template_xlsx() -> bytes:
        workbook = Workbook()
        instructions = workbook.active
        instructions.title = "Instrucciones"
        instructions.append(["PLANTILLA DE EVALUACIONES · AQU ANQA"])
        instructions.append([])
        for line in (
            "Complete una fila por evaluación de campo.",
            "module_key admite: estadios, flores, baya, pesos, brotes o ramas.",
            "Use lote_id o el conjunto fundo + modulo + lote.",
            "valores_json debe ser un objeto JSON con las métricas del módulo.",
            "Deje client_id vacío para que el previsualizador genere un UUID.",
            "La previsualización guarda un manifiesto, pero no escribe las tablas operativas.",
        ):
            instructions.append([line])
        instructions.column_dimensions["A"].width = 105

        sheet = workbook.create_sheet("Evaluaciones")
        sheet.append(TEMPLATE_HEADERS)
        sheet.append(TEMPLATE_EXAMPLE)
        header_fill = PatternFill("solid", fgColor="2D728F")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for index, header in enumerate(TEMPLATE_HEADERS, start=1):
            sheet.column_dimensions[chr(64 + index)].width = max(16, len(header) + 3)
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def preview_xlsx(content: bytes, filename: str) -> ImportPreview:
        if not content:
            raise ImportPreviewError("El archivo está vacío")
        if len(content) > 10 * 1024 * 1024:
            raise ImportPreviewError("El archivo supera el límite de 10 MB")
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as error:
            raise ImportPreviewError("No se pudo leer el libro Excel") from error
        sheet = (
            workbook["Evaluaciones"]
            if "Evaluaciones" in workbook.sheetnames
            else workbook.active
        )
        first_row = next(sheet.iter_rows(values_only=True), None)
        if first_row is None:
            raise ImportPreviewError("El libro no contiene una hoja con encabezados")
        headers = [str(value).strip() if value is not None else "" for value in first_row]
        positions = {header: index for index, header in enumerate(headers) if header}
        optional_headers = {"client_id", "captured_at", "fundo", "modulo", "lote"}
        missing = [
            header
            for header in TEMPLATE_HEADERS
            if header not in positions and header not in optional_headers
        ]
        if missing:
            raise ImportPreviewError(f"Faltan columnas: {', '.join(missing)}")

        rows: list[ImportPreviewRow] = []
        total = 0
        file_sha256 = hashlib.sha256(content).hexdigest()
        for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            raw = {
                header: values[index] if index < len(values) else None
                for header, index in positions.items()
            }
            if not any(value not in (None, "") for value in raw.values()):
                continue
            total += 1
            if total > 5000:
                raise ImportPreviewError("La previsualización admite como máximo 5.000 filas")
            try:
                payload = _payload_from_row(raw, file_sha256=file_sha256, row_number=excel_row)
                validated = EvaluationCreate.model_validate(payload)
                normalize_evaluation(validated)
                rows.append(
                    ImportPreviewRow(
                        row=excel_row,
                        valid=True,
                        payload=validated.model_dump(mode="json"),
                    )
                )
            except (ValidationError, ValueError, TypeError) as error:
                rows.append(
                    ImportPreviewRow(row=excel_row, valid=False, error=_error_message(error))
                )
        valid = sum(row.valid for row in rows)
        return ImportPreview(
            preview_id=uuid4(),
            file_sha256=file_sha256,
            filename=filename,
            sheet=sheet.title,
            total_rows=total,
            valid_rows=valid,
            invalid_rows=total - valid,
            rows=rows,
            message=(
                "Listo para revisión: las filas válidas cumplen el contrato de evaluación; "
                "todavía no se han escrito en PostgreSQL."
            ),
        )


def _payload_from_row(
    row: dict[str, Any], *, file_sha256: str, row_number: int
) -> dict[str, Any]:
    payload = {
        key: value
        for key, value in row.items()
        if key in TEMPLATE_HEADERS and value not in (None, "")
    }
    client_id = payload.get("client_id")
    payload["client_id"] = (
        str(client_id).strip()
        if client_id
        else str(uuid5(NAMESPACE_URL, f"aquanqa:admin-load:{file_sha256}:{row_number}"))
    )
    values = payload.get("valores_json", {})
    if isinstance(values, str):
        values = json.loads(values)
    if not isinstance(values, dict):
        raise ValueError("valores_json debe ser un objeto JSON")
    payload["valores"] = values
    payload.pop("valores_json", None)
    return payload


def _error_message(error: Exception) -> str:
    if isinstance(error, ValidationError):
        return "; ".join(
            f"{'.'.join(str(part) for part in item['loc'])}: {item['msg']}"
            for item in error.errors()
        )
    return str(error)


def _correction_values(
    module_key: ModuleKey,
    source_id: int,
    request: EvaluationCorrectionRequest,
) -> dict[str, Any]:
    """Valida y normaliza la edición antes de enviarla a la función SQL allowlisted."""
    if request.source_table == "ev_baya_medicion":
        diameter = request.valores.get("diametro")
        if diameter is None or isinstance(diameter, bool):
            raise ValueError("La medición de baya requiere un diámetro")
        suspicious = request.valores.get("sospechoso", False)
        if isinstance(suspicious, str):
            normalized = suspicious.strip().lower()
            if normalized not in {"true", "false"}:
                raise ValueError("sospechoso debe ser booleano")
            suspicious = normalized == "true"
        if not isinstance(suspicious, bool):
            raise ValueError("sospechoso debe ser booleano")
        return _json_safe_local(
            {
                "fecha": request.fecha,
                "lote_id": request.lote_id,
                "cortina": request.cortina,
                "hilera": request.hilera,
                "nro_muestra": request.nro_muestra,
                "diametro": diameter,
                "sospechoso": suspicious,
            }
        )

    payload = EvaluationCreate.model_validate(
        {
            "client_id": str(
                uuid5(NAMESPACE_URL, f"aquanqa:admin:{request.source_table}:{source_id}")
            ),
            "module_key": module_key,
            "fecha": request.fecha,
            "lote_id": request.lote_id,
            "cortina": request.cortina,
            "hilera": request.hilera,
            "planta": request.planta,
            "evaluador_id": request.evaluador_id,
            "item": request.item,
            "piso": request.piso,
            "hora": request.hora,
            "valores": request.valores,
        }
    )
    normalized = normalize_evaluation(payload)
    return _json_safe_local(
        {
            "fecha": payload.fecha,
            "lote_id": payload.lote_id,
            "cortina": payload.cortina,
            "hilera": payload.hilera,
            "planta": payload.planta,
            "evaluador_id": payload.evaluador_id,
            **normalized.data,
        }
    )


def _json_safe_local(value: Any) -> Any:
    """Convierte Decimal/time/date a JSON sin acoplar el caso de uso al adaptador SQL."""
    return json.loads(json.dumps(value, default=str))


__all__ = ["AdminService", "TEMPLATE_HEADERS", "ImportPreviewError"]
