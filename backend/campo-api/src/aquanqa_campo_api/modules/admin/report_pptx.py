"""Populate an Artifact Tool-authored native chart template with authorized report data."""
from copy import deepcopy
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

C='http://schemas.openxmlformats.org/drawingml/2006/chart'
A='http://schemas.openxmlformats.org/drawingml/2006/main'
P='http://schemas.openxmlformats.org/presentationml/2006/main'
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
for prefix,ns in [('c',C),('a',A),('p',P),('r',R)]: ET.register_namespace(prefix,ns)
COLORS=['145B45','23946F','6D82B8','D6913D','9477A9','358EAA','AC625B']
def xml(root):
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    if root.tag.startswith('{'+REL+'}') or root.tag.startswith('{'+CT+'}'):
        data=data.replace(b'ns0:',b'').replace(b':ns0',b'')
    return data
def sub(parent,ns,tag,attrs=None,text=None):
    node=ET.SubElement(parent,'{'+ns+'}'+tag,attrs or {})
    if text is not None:node.text=str(text)
    return node

def export_weekly_report(report):
    if not report.points or not report.desde or not report.hasta:raise ValueError('No hay datos para exportar')
    with ZipFile(Path(__file__).parent/'templates/weekly-report.pptx') as source:
        files={n:source.read(n) for n in source.namelist()}
    slide_template=ET.fromstring(files['ppt/slides/slide1.xml'])
    chart_template=ET.fromstring(files['ppt/slides/charts/chart1.xml'])
    slide_rels=ET.fromstring(files['ppt/slides/_rels/slide1.xml.rels'])
    presentation=ET.fromstring(files['ppt/presentation.xml'])
    ids=presentation.find('{'+P+'}sldIdLst');ids.clear()
    rels=ET.fromstring(files['ppt/_rels/presentation.xml.rels'])
    for node in list(rels):
        if node.get('Type','').endswith('/slide'):rels.remove(node)
    types=ET.fromstring(files['[Content_Types].xml'])
    for node in list(types):
        if node.get('PartName','') in ['/ppt/slides/slide1.xml','/ppt/slides/charts/chart1.xml']:types.remove(node)
    weeks=[];day=report.desde-timedelta(days=report.desde.weekday())
    while day<=report.hasta:weeks.append(day);day+=timedelta(days=7)
    labels=[f'S{d.isocalendar().week:02d}' for d in weeks]
    funds=sorted({p.fundo_id:p.fundo for p in report.points}.items())
    title='Flores observadas' if report.metric=='n_flores' else 'Cuajos observados'
    for index,(fund_id,fund) in enumerate(funds,1):
        points=[p for p in report.points if p.fundo_id==fund_id]
        modules=sorted({p.modulo_id:p.modulo for p in points}.items(), key=lambda item:item[1])
        by_key={(p.modulo_id,p.week):p for p in points}
        slide=deepcopy(slide_template)
        replacements={'{{TITLE}}':f'{title} · {fund}','{{PERIOD}}':f'{report.desde} a {report.hasta} · { {'registro_access':'Registro histórico','captura':'Capturas de campo'}.get(report.grano,report.grano) }',
            '{{BASE}}':f'{sum(p.available for p in points):,} con dato / {sum(p.evaluations for p in points):,} evaluaciones · Conteo por evaluación'}
        for text in slide.iter('{'+A+'}t'):
            for key,value in replacements.items():text.text=(text.text or '').replace(key,value)
        chart=deepcopy(chart_template);line=chart.find('.//{'+C+'}lineChart')
        prototype=deepcopy(line.find('{'+C+'}ser'))
        for old in list(line.findall('{'+C+'}ser')):line.remove(old)
        wb=Workbook();ws=wb.active;ws.title='Datos';ws.append(['Semana',*[name for _,name in modules]])
        for row,week in enumerate(weeks):ws.append([labels[row],*[by_key[(id,week)].mean if (id,week) in by_key else None for id,_ in modules]])
        for order,(module_id,name) in enumerate(modules):
            ser=deepcopy(prototype);ser.find('{'+C+'}idx').set('val',str(order));ser.find('{'+C+'}order').set('val',str(order))
            for color in ser.iter('{'+A+'}srgbClr'):color.set('val',COLORS[order%len(COLORS)])
            column=get_column_letter(order+2)
            tx=ser.find('{'+C+'}tx');tx.clear();ref=sub(tx,C,'strRef');sub(ref,C,'f',text=f"Datos!${column}$1")
            cache=sub(ref,C,'strCache');sub(cache,C,'ptCount',{'val':'1'});sub(sub(cache,C,'pt',{'idx':'0'}),C,'v',text=name)
            for tag,kind,formula,values in [('cat','str',f'Datos!$A$2:$A${len(weeks)+1}',labels),('val','num',f'Datos!${column}$2:${column}${len(weeks)+1}',[by_key[(module_id,w)].mean if (module_id,w) in by_key else None for w in weeks])]:
                node=ser.find('{'+C+'}'+tag);node.clear();ref=sub(node,C,kind+'Ref');sub(ref,C,'f',text=formula);cache=sub(ref,C,kind+'Cache')
                if kind=='num':sub(cache,C,'formatCode',text='0.0')
                sub(cache,C,'ptCount',{'val':str(len(values))})
                for i,value in enumerate(values):
                    if value is not None:sub(sub(cache,C,'pt',{'idx':str(i)}),C,'v',text=value)
            line.insert(2+order,ser)
        for axis in chart.findall('.//{'+C+'}valAx/{'+C+'}delete'):axis.set('val','0')
        for run in chart.iter('{'+A+'}defRPr'):run.set('sz','1200')
        category_axis=chart.find('.//{'+C+'}catAx')
        sub(category_axis,C,'tickLblSkip',{'val':str(max(1,(len(weeks)+25)//26))})
        external=sub(chart,C,'externalData',{'{'+R+'}id':'workbook'});sub(external,C,'autoUpdate',{'val':'0'})
        workbook=BytesIO();wb.save(workbook)
        files[f'ppt/embeddings/report{index}.xlsx']=workbook.getvalue()
        chartrels=ET.Element('{'+REL+'}Relationships');sub(chartrels,REL,'Relationship',{'Id':'workbook','Type':R+'/package','Target':f'../../embeddings/report{index}.xlsx'})
        files[f'ppt/slides/charts/_rels/chart{index}.xml.rels']=xml(chartrels)
        files[f'ppt/slides/charts/chart{index}.xml']=xml(chart)
        relationships=deepcopy(slide_rels)
        for node in list(relationships):
            if node.get('Type','').endswith('/chart'):node.set('Target',f'/ppt/slides/charts/chart{index}.xml')
            if node.get('Type','').endswith('/notesSlide'):relationships.remove(node)
        files[f'ppt/slides/slide{index}.xml']=xml(slide)
        files[f'ppt/slides/_rels/slide{index}.xml.rels']=xml(relationships)
        sub(ids,P,'sldId',{'id':str(255+index),'{'+R+'}id':f'reportSlide{index}'})
        sub(rels,REL,'Relationship',{'Id':f'reportSlide{index}','Type':R+'/slide','Target':f'/ppt/slides/slide{index}.xml'})
        sub(types,CT,'Override',{'PartName':f'/ppt/slides/slide{index}.xml','ContentType':'application/vnd.openxmlformats-officedocument.presentationml.slide+xml'})
        sub(types,CT,'Override',{'PartName':f'/ppt/slides/charts/chart{index}.xml','ContentType':'application/vnd.openxmlformats-officedocument.drawingml.chart+xml'})
    sub(types,CT,'Default',{'Extension':'xlsx','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'})
    files['ppt/presentation.xml']=xml(presentation);files['ppt/_rels/presentation.xml.rels']=xml(rels);files['[Content_Types].xml']=xml(types)
    output=BytesIO()
    with ZipFile(output,'w',ZIP_DEFLATED) as z:
        for name,data in files.items():z.writestr(name,data)
    return output.getvalue()
